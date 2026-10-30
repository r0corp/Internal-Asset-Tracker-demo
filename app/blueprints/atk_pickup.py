"""
Pengambilan ATK via scan QR rak - dioperasikan oleh Office Boy
(Sadam/Taufik/Roni), BUKAN self-service oleh pegawai. Alurnya:

1. Tiap kategori ATK punya QR sendiri (ditempel di rak), isinya link
   langsung ke /atk-pickup/new?category_id=<id>. QR ini cukup di-scan
   pakai kamera bawaan HP (bukan scanner khusus di dalam aplikasi) -
   produksi masih HTTP biasa (bukan HTTPS), jadi akses kamera browser
   (getUserMedia) tidak bisa diandalkan; buka link langsung jauh lebih
   sederhana dan tetap bekerja di semua HP.
2. Office Boy mengisi nama pegawai (wajib) + tanda tangan (opsional,
   salah satu dari keduanya cukup jadi bukti) + qty tiap item yang
   diambil (bisa lebih dari 1 kategori dalam 1 transaksi).
3. Submit langsung mengurangi Asset.stock_quantity saat itu juga.
4. Laporan (harian/mingguan/bulanan) dibaca dari riwayat transaksi ini,
   bisa diexport ke Excel & PDF.
"""

import io
import os
from datetime import date, datetime, timedelta

import qrcode
from flask import (
    render_template, request, redirect, url_for, flash, abort,
    send_file, current_app,
)
from flask_babel import gettext as _
from flask_login import current_user

from .. import db
from ..models import Asset, Category, AtkPickup, AtkPickupItem, FIELD_TEMPLATE_ATK
from ..auth import requires_permission

from . import main
from .helpers import apply_asset_category_restriction, log_action, parse_date


def _atk_categories():

    return (
        Category.query
        .filter_by(field_template=FIELD_TEMPLATE_ATK)
        .order_by(Category.stock_sequence.asc().nullslast(), Category.name.asc())
        .all()
    )


def _atk_assets():

    category_names = [c.name for c in _atk_categories()]

    return (
        apply_asset_category_restriction(
            Asset.query.filter(Asset.category.in_(category_names))
        )
        .order_by(Asset.category.asc(), Asset.asset_name.asc())
        .all()
    )


def _available_pcs(asset):
    """
    Stok yang tersedia dalam Pcs - pegawai selalu minta dalam Pcs,
    sedangkan Asset.stock_quantity dicatat dalam satuan aslinya
    (Box/Lusin/Pcs, lihat Asset.unit). Asset.pcs_per_unit = berapa Pcs
    isi 1 satuan tsb; kalau belum di-set, anggap 1 satuan = 1 Pcs
    (tidak ada konversi, sama seperti perilaku lama).
    """

    return round((asset.stock_quantity or 0) * (asset.pcs_per_unit or 1), 2)


def _fmt_qty(value):
    """Versi plain-Python dari filter Jinja `fmt_qty`, dipakai di export PDF."""

    if value is None:
        return "-"

    if value == int(value):
        return str(int(value))

    return f"{value:.2f}".rstrip("0").rstrip(".")

# ============================================================
# DAFTAR KATEGORI + QR (untuk dicetak & ditempel di rak)
# ============================================================


@main.route("/atk-pickup")
@requires_permission("atk_pickup", "read")
def atk_pickup_categories():

    categories = _atk_categories()

    return render_template(
        "atk_pickup/categories.html",
        categories=categories,
    )


@main.route("/atk-pickup/qr/<int:category_id>.png")
@requires_permission("atk_pickup", "read")
def atk_pickup_qr_image(category_id):

    category = Category.query.get_or_404(category_id)

    target_url = url_for(
        "main.atk_pickup_new",
        category_id=category.id,
        _external=True,
    )

    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=10,
        border=4,
    )
    qr.add_data(target_url)
    qr.make(fit=True)

    image = qr.make_image(fill_color="black", back_color="white")

    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)

    return send_file(buffer, mimetype="image/png")

# ============================================================
# FORM PENGAMBILAN
# ============================================================


@main.route(
    "/atk-pickup/new",
    methods=["GET", "POST"],
)
@requires_permission("atk_pickup", "create")
def atk_pickup_new():

    assets = _atk_assets()

    for asset in assets:
        asset.available_pcs = _available_pcs(asset)

    scanned_category_id = request.values.get(
        "category_id",
        type=int,
    )

    if request.method == "POST":

        employee_name = request.form.get(
            "employee_name",
            ""
        ).strip()

        signature = request.form.get(
            "signature",
            ""
        ).strip() or None

        if not employee_name:

            flash(
                _("Employee name is required."),
                "danger",
            )

            return redirect(
                url_for(
                    "main.atk_pickup_new",
                    category_id=scanned_category_id,
                )
            )

        rows = []
        errors = []

        for asset in assets:

            raw_qty = request.form.get(
                f"qty_{asset.id}",
                ""
            ).strip()

            if not raw_qty:
                continue

            try:
                qty = int(raw_qty)
            except ValueError:
                continue

            if qty <= 0:
                continue

            available_pcs = _available_pcs(asset)

            if qty > available_pcs:

                errors.append(
                    _(
                        "%(name)s: requested (%(qty)s Pcs) exceeds "
                        "available stock (%(stock)s Pcs)."
                    )
                    % {
                        "name": asset.asset_name,
                        "qty": qty,
                        "stock": int(available_pcs)
                        if available_pcs == int(available_pcs)
                        else round(available_pcs, 2),
                    }
                )

                continue

            rows.append((asset, qty))

        if errors:

            for message in errors[:10]:
                flash(message, "danger")

            return redirect(
                url_for(
                    "main.atk_pickup_new",
                    category_id=scanned_category_id,
                )
            )

        if not rows:

            flash(
                _("Select at least one item with a quantity."),
                "danger",
            )

            return redirect(
                url_for(
                    "main.atk_pickup_new",
                    category_id=scanned_category_id,
                )
            )

        pickup = AtkPickup(
            employee_name=employee_name,
            signature=signature,
            recorded_by=current_user.username,
        )
        db.session.add(pickup)
        db.session.flush()

        for asset, qty in rows:

            db.session.add(
                AtkPickupItem(
                    pickup_id=pickup.id,
                    asset_id=asset.id,
                    asset_name=asset.asset_name,
                    category=asset.category,
                    unit="Pcs",
                    quantity=qty,
                )
            )

            pcs_per_unit = asset.pcs_per_unit or 1

            asset.stock_quantity = round(
                (asset.stock_quantity or 0) - (qty / pcs_per_unit),
                4,
            )

        db.session.commit()

        log_action(
            "create",
            "AtkPickup",
            entity_label=f"{employee_name} ({len(rows)} item)",
        )

        flash(
            _("Pickup for %(name)s saved successfully.")
            % {"name": employee_name},
            "success",
        )

        return redirect(
            url_for("main.atk_pickup_new")
        )

    category_id_by_name = {
        c.name: c.id
        for c in _atk_categories()
    }

    return render_template(
        "atk_pickup/new.html",
        assets=assets,
        scanned_category_id=scanned_category_id,
        category_id_by_name=category_id_by_name,
    )

# ============================================================
# LAPORAN
# ============================================================


def _period_range(period, anchor):

    if period == "daily":
        start = anchor
        end = anchor + timedelta(days=1)

    elif period == "monthly":
        start = anchor.replace(day=1)
        if start.month == 12:
            end = start.replace(year=start.year + 1, month=1)
        else:
            end = start.replace(month=start.month + 1)

    else:
        # weekly (default) - Senin s/d Minggu minggu yang sama.
        start = anchor - timedelta(days=anchor.weekday())
        end = start + timedelta(days=7)

    return start, end


def _report_data(period, anchor):

    start, end = _period_range(period, anchor)

    pickups = (
        AtkPickup.query
        .filter(
            AtkPickup.created_at >= datetime.combine(start, datetime.min.time()),
            AtkPickup.created_at < datetime.combine(end, datetime.min.time()),
        )
        .order_by(AtkPickup.created_at.desc())
        .all()
    )

    totals = {}

    for pickup in pickups:
        for item in pickup.items:

            key = item.asset_name

            if key not in totals:

                totals[key] = {
                    "asset_name": item.asset_name,
                    "category": item.category,
                    "unit": item.unit,
                    "qty": 0,
                    "remaining": (
                        _available_pcs(item.asset)
                        if item.asset else None
                    ),
                }

            totals[key]["qty"] += item.quantity

    totals_list = sorted(
        totals.values(),
        key=lambda row: (row["category"] or "", row["asset_name"]),
    )

    return start, end, pickups, totals_list


@main.route("/atk-pickup/report")
@requires_permission("atk_pickup", "read")
def atk_pickup_report():

    period = request.args.get("period", "weekly")

    if period not in ("daily", "weekly", "monthly"):
        period = "weekly"

    anchor = parse_date(request.args.get("date")) or date.today()

    start, end, pickups, totals = _report_data(period, anchor)

    return render_template(
        "atk_pickup/report.html",
        period=period,
        anchor=anchor.strftime("%Y-%m-%d"),
        start=start,
        end=end - timedelta(days=1),
        pickups=pickups,
        totals=totals,
    )


@main.route("/atk-pickup/report/export.xlsx")
@requires_permission("atk_pickup", "read")
def atk_pickup_export_excel():

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    period = request.args.get("period", "weekly")

    if period not in ("daily", "weekly", "monthly"):
        period = "weekly"

    anchor = parse_date(request.args.get("date")) or date.today()

    start, end, pickups, totals = _report_data(period, anchor)
    end_display = end - timedelta(days=1)

    wb = Workbook()

    # ------------------------------------------------------------
    # SHEET 1 - REKAP PER ITEM
    # ------------------------------------------------------------

    ws1 = wb.active
    ws1.title = "Rekap"

    header_fill = PatternFill(
        start_color="0D6EFD", end_color="0D6EFD", fill_type="solid"
    )
    header_font = Font(color="FFFFFF", bold=True)

    ws1.append([
        f"ATK Pickup Recap ({start.strftime('%d-%m-%Y')} - {end_display.strftime('%d-%m-%Y')})"
    ])
    ws1.append([])
    ws1.append(["Category", "Item", "Unit", "Qty Taken", "Remaining Stock"])

    for cell in ws1[3]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in totals:
        ws1.append([
            row["category"] or "-",
            row["asset_name"],
            row["unit"] or "-",
            row["qty"],
            row["remaining"] if row["remaining"] is not None else "-",
        ])

    for col, width in zip("ABCDE", (20, 30, 10, 12, 16)):
        ws1.column_dimensions[col].width = width

    # ------------------------------------------------------------
    # SHEET 2 - RIWAYAT TRANSAKSI
    # ------------------------------------------------------------

    ws2 = wb.create_sheet("History")
    ws2.append(["Date/Time", "Employee", "Item", "Qty", "Recorded By", "Signed"])

    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for pickup in pickups:
        for item in pickup.items:
            ws2.append([
                pickup.created_at.strftime("%d-%m-%Y %H:%M"),
                pickup.employee_name,
                item.asset_name,
                item.quantity,
                pickup.recorded_by or "-",
                "Yes" if pickup.signature else "No",
            ])

    for col, width in zip("ABCDEF", (18, 20, 30, 8, 14, 10)):
        ws2.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"atk-pickup-{period}-{anchor.strftime('%Y%m%d')}.xlsx"

    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


@main.route("/atk-pickup/report/export.pdf")
@requires_permission("atk_pickup", "read")
def atk_pickup_export_pdf():

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    period = request.args.get("period", "weekly")

    if period not in ("daily", "weekly", "monthly"):
        period = "weekly"

    anchor = parse_date(request.args.get("date")) or date.today()

    start, end, pickups, totals = _report_data(period, anchor)
    end_display = end - timedelta(days=1)

    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()
    story = []

    story.append(
        Paragraph(
            f"ATK Pickup Recap - {start.strftime('%d-%m-%Y')} to {end_display.strftime('%d-%m-%Y')}",
            styles["Title"],
        )
    )
    story.append(Spacer(1, 10))

    recap_data = [["Category", "Item", "Unit", "Qty Taken", "Remaining Stock"]]
    for row in totals:
        recap_data.append([
            row["category"] or "-",
            row["asset_name"],
            row["unit"] or "-",
            str(row["qty"]),
            _fmt_qty(row["remaining"]),
        ])

    recap_table = Table(recap_data, repeatRows=1)
    recap_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d6efd")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(recap_table)
    story.append(Spacer(1, 20))

    story.append(Paragraph("Pickup History", styles["Heading2"]))
    story.append(Spacer(1, 6))

    history_data = [["Date/Time", "Employee", "Item", "Qty", "Recorded By", "Signed"]]
    for pickup in pickups:
        for item in pickup.items:
            history_data.append([
                pickup.created_at.strftime("%d-%m-%Y %H:%M"),
                pickup.employee_name,
                item.asset_name,
                str(item.quantity),
                pickup.recorded_by or "-",
                "Yes" if pickup.signature else "No",
            ])

    history_table = Table(history_data, repeatRows=1)
    history_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d6efd")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(history_table)

    doc.build(story)
    buffer.seek(0)

    filename = f"atk-pickup-{period}-{anchor.strftime('%Y%m%d')}.pdf"

    return send_file(
        buffer,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=filename,
    )
